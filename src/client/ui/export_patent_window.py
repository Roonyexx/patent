import os
import tkinter as tk
from tkinter import messagebox, filedialog
from docx import Document


WINDOW_SIZE = '500x120'


class ExportPatentWindow:
    def __init__(self, parent, values):
        self.values = values

        self.dialog = tk.Toplevel(parent)
        self.dialog.title("Экспорт")
        self.dialog.geometry(WINDOW_SIZE)
        self.dialog.transient(parent)
        self.dialog.grab_set()

        self.path_var = tk.StringVar(self.dialog)
        self.file_name_var = tk.StringVar(self.dialog)

        self.create_widgets()

    def create_widgets(self):
        frame = tk.Frame(self.dialog, pady=10, padx=10)
        frame.pack(fill='both', expand=True)

        for i in range(3):
            frame.rowconfigure(index=i, weight=1)

        for j in range(2):
            frame.columnconfigure(index=j, weight=1)

        tk.Label(frame, text="Путь:").grid(row=0, column=0, sticky='w')

        path_frame = tk.Frame(frame)
        path_frame.grid(row=0, column=1, sticky='ew')
        path_frame.columnconfigure(index=0, weight=2)
        path_frame.columnconfigure(index=1, weight=1)

        self.path_var.set(os.path.join(os.getcwd()))
        tk.Entry(path_frame, textvariable=self.path_var, state='disabled').grid(row=0, column=0, sticky='ew')

        tk.Button(path_frame, text="Обзор...", command=self.browse_path).grid(row=0, column=1)

        tk.Label(frame, text="Имя файла:").grid(row=1, column=0, sticky='w')
        tk.Entry(frame, textvariable=self.file_name_var).grid(row=1, column=1, sticky='ew')

        tk.Button(frame, text="Экспортировать отчет", command=self.export_report, width=25).grid(row=2, column=0)

    def browse_path(self):
        filename = filedialog.askdirectory(
            initialdir=os.getcwd()
        )

        if filename:
            self.path_var.set(filename)

    def export_report(self):
        try:
            if not self.values:
                messagebox.showwarning("Предупреждение", "Нет параметров для сохранения!")
                return

            filepath = self.path_var.get()

            if not filepath:
                messagebox.showerror('Ошибка', 'Укажите путь для сохранения файла')
                return

            if not self.file_name_var.get():
                messagebox.showerror('Ошибка', 'Укажите имя файла')
                return

            doc = Document()
            doc.add_heading(f'Патент №{self.values.get('id')}', 0)

            for key, value in self.values.items():
                para = doc.add_paragraph()
                run_key = para.add_run(f"{key}: ")
                run_key.bold = True
                para.add_run(str(value))

            filepath += '/' + self.file_name_var.get() + '.doc'
            doc.save(filepath)
            result = messagebox.askyesno(
                'Успешный экспорт',
                f'Документ успешно сохранен:\n{filepath}\n\nОткрыть папку с файлом?'
            )

            if result:
                os.startfile(os.path.dirname(filepath))
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось создать документ:\n{str(e)}")

    def show(self):
        self.dialog.protocol("WM_DELETE_WINDOW", self.on_closing)
        self.dialog.mainloop()

    def on_closing(self):
        self.dialog.destroy()
